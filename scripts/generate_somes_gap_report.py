#!/usr/bin/env python3
"""Generate the UC3.3 SoMES vs UC3.2 gap/delegation Word report."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


NAVY = RGBColor(0x1B, 0x3A, 0x4B)
ACCENT = RGBColor(0x1F, 0x5F, 0x74)
MUTED = RGBColor(0x4A, 0x55, 0x5C)
GREEN = RGBColor(0x1B, 0x6B, 0x3A)
AMBER = RGBColor(0x9A, 0x6B, 0x12)
RED = RGBColor(0x8B, 0x1E, 0x1E)


def set_run_font(run, *, name="Calibri", size=11, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "C5CDD1")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5F74")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rPr.append(sz)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rFonts)
    new_run.append(rPr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def style_heading(paragraph, level: int) -> None:
    for run in paragraph.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
        if level == 1:
            run.font.size = Pt(16)
        elif level == 2:
            run.font.size = Pt(13)
        else:
            run.font.size = Pt(12)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    style_heading(p, level)
    return p


def add_para(doc: Document, text: str, *, bold=False, italic=False, size=11, space_after=8, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color or MUTED)
    return p


def add_bullet(doc: Document, text: str, *, level=0, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(1.0 + 0.6 * level)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True, color=NAVY)
        r2 = p.add_run(text)
        set_run_font(r2, size=11, color=MUTED)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=MUTED)
    return p


def add_table(doc: Document, headers: list[str], rows: list[list[str]], col_widths: list[float] | None = None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(hdr[i], "1B3A4B")
        set_cell_border(hdr[i])
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        bg = "F4F7F8" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            run = p.add_run(val)
            color = MUTED
            if val.startswith("Hotové") or val.startswith("Splnené") or val.startswith("Áno"):
                color = GREEN
            elif val.startswith("Čiastočné") or val.startswith("Externá") or val.startswith("Partial") or "závislosť" in val:
                color = AMBER
            elif val.startswith("Chýba") or val.startswith("Nie") or val.startswith("Otvorené"):
                color = RED
            set_run_font(run, size=9.5, color=color, bold=c_idx == 0)
            shade_cell(cells[c_idx], bg)
            set_cell_border(cells[c_idx])
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("UC3 Smart Grid  ·  SoMES / SEED  ·  interný podklad pre nábor  ·  dôverné")
    set_run_font(r, size=8, color=RGBColor(0x7A, 0x86, 0x8C), italic=True)

    # Header block
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.LEFT
    kr = kicker.add_run("UC3  ·  SMART GRID MANAGEMENT  ·  SPICE / DOMINO")
    set_run_font(kr, size=10, bold=True, color=ACCENT)
    kicker.paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph()
    tr = title.add_run("Migrácia UC3.3 do SoMES — gap analýza a návrh práce pre dvoch programátorov")
    set_run_font(tr, size=22, bold=True, color=NAVY)
    title.paragraph_format.space_after = Pt(8)

    meta = doc.add_paragraph()
    mr = meta.add_run(
        "Stav k 27. augustu 2026  ·  porovnanie kódu v GitHub s referenčnou architektúrou SoMES "
        "a s migračným zadaním  ·  určené na onboarding a delegovanie práce"
    )
    set_run_font(mr, size=11, italic=True, color=MUTED)
    meta.paragraph_format.space_after = Pt(14)

    add_para(
        doc,
        "Tento dokument je interný podklad pre prijatie dvoch programátorov. "
        "Nie je to architektúrny verdikt namiesto existujúcej implementačnej správy. "
        "Je to porovnanie toho, čo je v kóde skutočne hotové, čo ostáva otvorené, "
        "a akú prácu je rozumné na nových ľudí delegovať bez toho, aby sa rozbila "
        "už postavená SoMES kostra.",
        size=11,
        space_after=14,
    )

    # ---------------------------------------------------------------------
    add_heading(doc, "1. Zhrnutie v bodoch", 1)

    add_heading(doc, "1.1 Verdikt", 2)
    add_bullet(doc, "Migrácia UC3.3 → SoMES je architektonicky úspešná. Referenčných 18 modulov má v kóde pokrytie; jadro je operačné, nie investičné.")
    add_bullet(doc, "Implementácia v scditech/uc3_3_somes (v0.2.7) zodpovedá záveru implementačnej správy: cca 90–95 % súlad s cieľovou SoMES architektúrou.")
    add_bullet(doc, "Hlavný cieľ migrácie je splnený: výstupom je next-day dispatch plán v 15-min kroku, technická validácia, KPI, dashboard a EMS/BEMS payload — nie payback/NPV.")
    add_bullet(doc, "Investičná logika (sizing, CAPEX, NPV, payback, katalógy) ostáva v UC3.2 industry_sg_vre_cost_optimizer (v0.1.42). SoMES ju nepočíta; dashboard ju berie len ako SEED appendix.")
    add_bullet(doc, "To, čo ostáva, nie je nová architektúra. Je to produkčné doriešenie: reálne merania, EMS endpoint, testy, zatvorenie feedback slučky, silnejšia grid validácia a upratanie leftover pieces.")

    add_heading(doc, "1.2 Čo je hotové (nechať, neprepisovať)", 2)
    add_bullet(doc, "Data ingestion, validácia kvality dát, preprocessing a load forecast (XGBoost).")
    add_bullet(doc, "PV forecast cez Open-Meteo + UC3.4 PVOUT pipeline; price forecast cez OKTE + PriceForecastPiece.")
    add_bullet(doc, "Dispatch optimalizácia ako LP (SciPy/HiGHS) v common_somes/dispatch.py, volaná z BatterySimPiece; greedy fallback pri zlyhaní solvera.")
    add_bullet(doc, "Battery plan, grid import/export, flexible load schedule, next-day dispatch CSV.")
    add_bullet(doc, "Technická validácia (batéria, menič, prípojka) v GridFeasibilityPiece.")
    add_bullet(doc, "Operačné KPI, dashboard (JSON + HTML), forecast-vs-actual v ModelMonitoringPiece.")
    add_bullet(doc, "EMS klient: reálne HTTP POST, auth, retry, ACK, Modbus register mapa (súbor, nie live protokol).")
    add_bullet(doc, "SomesConnectorsPiece s provenance a poradím zdrojov: súbor → URL → live API → demo fixture.")

    add_heading(doc, "1.3 Čo ešte nie je dorobené", 2)
    add_bullet(doc, "Reálny PV merací feed — architektúra pripravená, v default behu syntetika / Open-Meteo PVOUT.", bold_prefix="Externá závislosť.  ")
    add_bullet(doc, "Reálna BESS telemetria (BMS) — default je demo fixture.", bold_prefix="Externá závislosť.  ")
    add_bullet(doc, "Produkčný EMS/BEMS endpoint, credentials, end-to-end ACK na reálnom zariadení.", bold_prefix="Externá závislosť.  ")
    add_bullet(doc, "Grid Feasibility používa linearizovaný odhad napätia na PCC, nie full network power-flow. Jediné architektonicky čiastočné miesto.", bold_prefix="Architektúra.  ")
    add_bullet(doc, "Flexible loads sú greedy (najlacnejšie okno), nie spoločná optimalizácia s batériou v LP.", bold_prefix="Kvalita algoritmu.  ")
    add_bullet(doc, "Forecast-vs-actual dáva odporúčania, ale IncrementalTrainPiece nie je automaticky spúšťaný. Slučka je advisory, nie closed-loop.", bold_prefix="Closed loop.  ")
    add_bullet(doc, "Takmer žiadne testy dispatch/EMS/grid/connectors (2 test súbory, ~13 testov, len Open-Meteo a preprocessing).", bold_prefix="Kvalita kódu.  ")
    add_bullet(doc, "V pieces/ ostáva 6 orphan pieces mimo Domino importu (SolarSim, Simulate, Evaluate, Explainable, ForecastAggregator, SyntheticData).", bold_prefix="Upratanie.  ")
    add_bullet(doc, "UC3.2 nemá FeasibilityReportPiece, nemá seed_handoff.json, nemá testy, v kóde je natvrdo OneData token.", bold_prefix="SEED strana.  ")

    add_heading(doc, "1.4 Čo dať dvom programátorom", 2)
    add_bullet(
        doc,
        "Integrácie a prevádzka SoMES. Reálne konektory (PV, BESS, EMS), secrets, OneData produkčné cesty, zatvorenie feedback slučky, dashboard/alerty, testy konektorov a EMS klienta. Necháva architektúru; doriešuje produkciu.",
        bold_prefix="Programátor A — SoMES production.  ",
    )
    add_bullet(
        doc,
        "Algoritmy, validácia a SEED kontrakt. Testy LP dispatchu, spoluoptimalizácia flexible loads, voliteľný pandapower power-flow, upratanie leftover pieces, FeasibilityReport / seed_handoff na UC3.2, odstránenie duplicít SolarSim/BatterySim/Simulate medzi repositármi.",
        bold_prefix="Programátor B — optimalizácia + SEED.  ",
    )
    add_bullet(doc, "Nezadávať im prepis SoMES architektúry, nový solver „od nuly“, ani investičný dashboard ako hlavný výstup SoMES. To by vrátilo workflow do stavu pred migráciou.")

    # ---------------------------------------------------------------------
    add_heading(doc, "2. Zdroje a rozsah porovnania", 1)
    add_para(doc, "Porovnanie vychádza z kódu v GitHub k dátumu tohto dokumentu a z migračného zadania, ktoré bolo dodané ako vstup. Google dokumenty referenčnej architektúry a implementačnej správy nie sú verejne čitateľné; ich obsah je zohľadnený v rozsahu, v akom bol vložený do zadania.")

    add_heading(doc, "2.1 Repositáre", 2)
    add_table(
        doc,
        ["Repositár", "Rola", "Verzia", "Čo to je"],
        [
            [
                "scditech/uc3_3_somes",
                "SoMES (operačný)",
                "0.2.7",
                "gUC3_3_SoMES. Next-day dispatch, forecast, validácia, EMS, dashboard. Domino import 26 nodov.",
            ],
            [
                "filipchrvala/industry_sg_vre_cost_optimizer",
                "SEED / UC3.2 (investičný)",
                "0.1.42",
                "Pôvodný MRK cost optimizer. Sizing, katalógy, CAPEX, NPV, payback, CFO dashboard.",
            ],
        ],
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SoMES: ")
    set_run_font(r, bold=True, color=NAVY)
    add_hyperlink(p, "https://github.com/scditech/uc3_3_somes", "https://github.com/scditech/uc3_3_somes")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("UC3.2 / SEED: ")
    set_run_font(r, bold=True, color=NAVY)
    add_hyperlink(p, "https://github.com/filipchrvala/industry_sg_vre_cost_optimizer", "https://github.com/filipchrvala/industry_sg_vre_cost_optimizer")

    add_heading(doc, "2.2 Zadanie, voči ktorému sa meria", 2)
    add_bullet(doc, "SoMES Workflow Type Reference Architecture (18 logických modulov).")
    add_bullet(doc, "Migration of UC3.3 SoMES workflow type — keep / refactor / move / add.")
    add_bullet(doc, "Interné vyhodnotenie migrácie (coverage matrix, 90–95 % súlad, otvorené externé závislosti).")
    add_bullet(doc, "Filipova implementačná správa UC33_SoMES_implementacia (záver o 21 hotových, 3 čakajúcich, 1 čiastočnom module).")

    add_heading(doc, "2.3 Čo dokument zámerne nerobí", 2)
    add_bullet(doc, "Neprehlasuje implementačnú správu za neplatnú. Kód ju v zásade potvrdzuje.")
    add_bullet(doc, "Nepočíta personodni. Prácu člení podľa zložky systému a rizika, nie podľa kalendára.")
    add_bullet(doc, "Nedáva juniorom úlohy, ktoré vyžadujú zmenu referenčnej architektúry bez architekta.")

    # ---------------------------------------------------------------------
    add_heading(doc, "3. Dva workflowy vedľa seba", 1)

    add_heading(doc, "3.1 UC3.2 Cost Optimizer — investičný reťazec", 2)
    add_para(
        doc,
        "Dvanásť Domino pieces. Jadro je SimulatePiece (~2 100 riadkov): MRK simulácia, CAPEX, NPV, payback. "
        "Ostatné pieces sú tenké obaly nad tým istým engine.",
    )
    add_para(doc, "Reťazec:", bold=True, space_after=4)
    add_para(
        doc,
        "UserInput → CatalogSync → TechnicalLimits → SizingOptimization → CatalogRanker → "
        "SolarSim → BatteryStrategy → BatterySim → Simulate → KPI / InvestmentEval → Dashboard",
        italic=True,
    )
    add_para(
        doc,
        "Toto je správne miesto pre investičné rozhodnutie. Nie je to SoMES. "
        "Problém je, že tu ostávajú aj operačné kópie SolarSim / BatterySim / Simulate, "
        "ktoré SoMES už má vo vlastnej, dispatch-oriented podobe. Bez kontraktu medzi repositármi "
        "hrozí, že sa budú rozchádzať.",
    )

    add_heading(doc, "3.2 UC3.3 SoMES — operačný reťazec", 2)
    add_para(
        doc,
        "Domino import (uc33_somes.customization) má 26 nodov a 47 hrán. "
        "README hovorí o 18-nodovom DAG; v importe je viac, lebo PV vetva je rozbitá na UC3.4 pieces. "
        "To sedí s poznámkou implementačnej správy: 25 implementačných modulov voči 18 logickým.",
    )
    add_para(doc, "Tri vetvy, ktoré sa zídu na dispatch:", bold=True, space_after=4)
    add_bullet(doc, "Load: FetchEnergyData → Preprocess → TrainModel → Predict → ModelMonitoring")
    add_bullet(doc, "PV: OpenMeteoPVData → DataPreprocessing → ModelDecider → DataNormalization → PvoutFeatureSelect → PVOUT train → error correction → staged spec → Inference → PvoutToVirtualSolar")
    add_bullet(doc, "Prevádzka: SomesConnectors → PriceForecast; IncrementalTrain → ForecastHorizon → FlexibleLoadSchedule; BatteryStrategy → BatterySim → GridFeasibility → EmsBemsOutput → Dashboard")
    add_para(
        doc,
        "Investičné pieces (TechnicalLimits, SizingOptimization, InvestmentEval, Catalog*, UserInput, FeasibilityReport) "
        "v SoMES repositári nie sú. To je správny výsledok migrácie.",
    )

    add_heading(doc, "3.3 Mapovanie keep / refactor / move", 2)
    add_table(
        doc,
        ["Pôvodné UC3.3", "Rozhodnutie v zadaní", "Stav v kóde"],
        [
            ["FetchEnergyData, Preprocess", "Keep and refactor", "Hotové — v SoMES DAG"],
            ["Train / Predict / ForecastHorizon", "Keep ako load forecast", "Hotové — v SoMES DAG"],
            ["ModelMonitoring, Dashboard", "Keep and adapt", "Hotové — operačný dashboard, nie CFO"],
            ["SolarSim", "Refactor strongly na PV forecast", "Hotové inou cestou (Open-Meteo + PVOUT); SolarSimPiece je leftover mimo importu"],
            ["BatteryStrategy + BatterySim", "Refactor strongly na dispatch", "Hotové — LP HiGHS v BatterySim; názov piece ostal „simulation“"],
            ["Simulate", "Refactor, nie lift-and-shift", "Mimo importu; ostáva ako knižnica pre legacy sim"],
            ["TechnicalLimits, Sizing", "Move mimo SoMES / SEED", "Splnené — v SoMES nie sú, ostávajú v UC3.2"],
            ["InvestmentEval, FeasibilityReport", "Move do SEED", "Z SoMES odstránené; FeasibilityReport v UC3.2 stále chýba"],
            ["WebUserInput", "Len configuration adapter", "V SoMES nie je; vstupy idú cez OneData / scenario.yaml"],
            ["AnomalyAlert, SustainableIngest", "Keep ako monitoring / ingest", "Ako samostatné pieces v SoMES nie sú; funkcia je v connectors / feedback / IncrementalTrain"],
        ],
    )

    # ---------------------------------------------------------------------
    add_heading(doc, "4. Coverage voči 18 modulom SoMES", 1)
    add_para(
        doc,
        "Tabuľka meria kód v scditech/uc3_3_somes, nie deklaráciu v dokumentácii. "
        "Stav „Hotové“ znamená, že modul má implementáciu a je v Domino importe. "
        "Neznamená to, že má produkčný dátový zdroj.",
    )
    add_table(
        doc,
        ["#", "SoMES modul", "Kde v kóde", "Stav"],
        [
            ["1", "Data Ingestion", "SomesConnectorsPiece, FetchEnergyDataPiece", "Hotové"],
            ["2", "Data Validation", "common_somes/quality.py + connectors", "Hotové"],
            ["3", "Data Preprocessing", "PreprocessEnergyDataPiece, DataPreprocessingPiece", "Hotové"],
            ["4", "Load Forecast", "TrainModel, Predict, ForecastHorizon, IncrementalTrain", "Hotové"],
            ["5", "PV / RES Forecast", "OpenMeteoPVData → PVOUT train/inference → PvoutToVirtualSolar", "Hotové (label je syntetické PVOUT z irradiance, nie meraný PV výkon)"],
            ["6", "Price / Tariff Forecast", "PriceForecastPiece + OKTE v connectors", "Hotové"],
            ["7", "Dispatch Optimization", "common_somes/dispatch.py (HiGHS LP) cez BatterySimPiece", "Hotové"],
            ["8", "Battery Charge/Discharge Plan", "výstup BatterySimPiece / next_day_dispatch_plan", "Hotové"],
            ["9", "Grid Import / Export Plan", "rovnaký dispatch výstup (import/export stĺpce)", "Hotové"],
            ["10", "Flexible Load Schedule", "FlexibleLoadSchedulePiece — greedy, nie LP", "Čiastočné"],
            ["11", "Battery Constraints Check", "GridFeasibilityPiece / architecture.py", "Hotové"],
            ["12", "Inverter & Connection Limits", "GridFeasibilityPiece", "Hotové"],
            ["13", "Grid Feasibility Check", "linearizovaný PCC dU/Un = (P·R + Q·X) / Un²", "Čiastočné"],
            ["14", "Next-Day Dispatch Plan", "approved_next_day_operating_plan.csv po validácii", "Hotové"],
            ["15", "Operational KPIs", "architecture.operational_kpis_from_dispatch (nie samostatný KPIPiece)", "Hotové"],
            ["16", "Monitoring Dashboard", "DashboardPiece — JSON + HTML", "Hotové"],
            ["17", "Output API / EMS-BEMS", "EmsBemsOutputPiece + ems_client.py", "Hotové (funkčné, bez produkčného endpointu)"],
            ["18", "Forecast vs Actual", "ModelMonitoringPiece + feedback.py", "Hotové (advisory, nie automatický retrain)"],
        ],
    )

    add_heading(doc, "4.1 Enabling pieces, ktoré architektúra nespomína", 2)
    add_para(
        doc,
        "Počet pieces je vyšší ako 18, lebo implementácia rozdelila logické moduly. "
        "To nie je chyba. Odporúčané mapovanie zo zadania ostáva v platnosti:",
    )
    add_bullet(doc, "Data Ingestion  →  SomesConnectorsPiece")
    add_bullet(doc, "Price / Tariff Forecast  →  PriceForecastPiece")
    add_bullet(doc, "Output API / EMS-BEMS  →  EmsBemsOutputPiece")
    add_para(
        doc,
        "Tieto tri pieces by sa v UC3 Architecture Document nemali vystavovať ako nové architektonické moduly. "
        "Sú to implementačné komponenty logických modulov 1, 6 a 17.",
    )

    add_heading(doc, "4.2 Pieces mimo Domino importu (leftover)", 2)
    add_table(
        doc,
        ["Piece", "Prečo tam je", "Odporúčanie"],
        [
            ["SolarSimPiece", "Pôvodný syntetický PV profil", "Nepoužívať v SoMES DAG; buď zmazať, alebo nechať ako fallback mimo importu"],
            ["SimulatePiece", "MRK investičná simulácia", "Nesmie byť v SoMES DAG; ak treba knižnicu, presunúť do shared / UC3.2"],
            ["EvaluateMLModelPiece", "UC3.4 template, no-op bez payloadu", "Buď zapojiť do PV vetvy, alebo vyradiť z produkčného obrazu"],
            ["ExplainablePredictionPiece", "UC3.4 template", "Nice-to-have, nie súčasť 18 modulov"],
            ["ForecastAggregatorPiece", "Z importu už vypadlo", "Nechať mimo, kým nebude multi-model požiadavka"],
            ["SyntheticDataGeneratorPiece", "Demo dáta", "Dev tool, nie produkčný modul"],
        ],
    )

    # ---------------------------------------------------------------------
    add_heading(doc, "5. Detail: čo ostáva otvorené", 1)

    add_heading(doc, "5.1 Externé závislosti (bez partnera / site dát sa nedokončia)", 2)
    add_para(
        doc,
        "Tieto tri body implementačná správa správne označuje ako externé. "
        "Programátor ich vie pripraviť (kontrakt, adapter, test s fixture), "
        "ale nemôže ich uzavrieť bez reálneho zdroja.",
    )
    add_table(
        doc,
        ["Položka", "Čo kód už vie", "Čo chýba od vonku", "Čo vie doprogramovať nový človek"],
        [
            [
                "PV meranie",
                "connectors: file / URL / demo; PV forecast z Open-Meteo",
                "Produkčný feed meraného PV výkonu (kW) v 15 min",
                "Adapter, schema, quality check, prepnutie z demo na live, testy",
            ],
            [
                "BESS telemetria",
                "bess_telemetry_url + demo SOC/power fixture; initial SOC ide do LP",
                "BMS/EMS telemetria (SOC, limity, dostupnosť)",
                "Normalizácia BMS CSV/API, mapovanie bodov, fail-soft keď BMS vypadne",
            ],
            [
                "EMS endpoint",
                "HTTP POST, bearer/basic/api_key, retry, ACK parser, Modbus mapa CSV",
                "URL, credentials, dohodnutý payload, ACK schéma cieľového EMS",
                "Konfigurácia, staging mock EMS, end-to-end test, runbook keď ACK nepríde",
            ],
        ],
    )

    add_heading(doc, "5.2 Jediné architektonicky čiastočné miesto: Grid Feasibility", 2)
    add_para(
        doc,
        "Kód to priznáva explicitne: nie je to load-flow downstream siete. "
        "Je to linearized voltage approximation v bode pripojenia. "
        "Kontroluje aj import/export limity, inverter limit a C-rate batérie. "
        "Pri porušení plán orezáva (clipping), nie re-dispatch.",
    )
    add_bullet(doc, "Chýba: sieťová topológia, pandapower / OpenDSS, Q z load/PV, nadväznosť na flexible loads, re-optimalizácia po reject.")
    add_bullet(doc, "Odporúčanie: nenechať juniora stavať full power-flow ako prvú úlohu. Najprv testy existujúcej validácie a korekčného plánu. Full network model až keď je site model (impedancie, transformátor, limity) k dispozícii.")

    add_heading(doc, "5.3 Inžinierske medzery, ktoré správa o 90–95 % podceňuje", 2)
    add_para(
        doc,
        "Architektúra je hotová. Produkčná tvrdosť nie. Toto je práca, ktorú dvom programátorom "
        "dáme prednostne — je delegovateľná, merateľná a neničí migráciu.",
    )

    add_para(doc, "Flexible loads nie sú v objektíve LP.", bold=True, space_after=4)
    add_para(
        doc,
        "FlexibleLoadSchedulePiece vyberie najlacnejšie časové okno v rámci time bandu. "
        "BatterySimPiece následne pripočíta flexible_load_kw k záťaži. "
        "Batéria teda vidí už posunutú záťaž, ale nerozhoduje o shifte spoločne s nabíjaním. "
        "Pre demo to stačí. Pre priemyselný microgrid s HVAC a batch procesom to nie je optimum.",
    )

    add_para(doc, "Feedback loop je otvorená.", bold=True, space_after=4)
    add_para(
        doc,
        "ModelMonitoringPiece počíta MAE/RMSE/MAPE pre load, PV aj dodržanie dispatchu. "
        "Vie povedať „spusti IncrementalTrainPiece“. IncrementalTrainPiece je však samostatný source node, "
        "nie automatický follow-up. closed_loop_ok je flag, nie orchestrácia. "
        "Dva BatterySim nody v DAG to ešte komplikujú: preview (node 110) ide do monitoringu na historickom loade, "
        "produkčný D+1 plán (node 114) ide do Grid/EMS. Porovnanie plán vs skutočnosť teda nemusí meriať ten istý plán, ktorý odišiel do EMS.",
    )

    add_para(doc, "Testy takmer nie sú.", bold=True, space_after=4)
    add_para(
        doc,
        "SoMES: test_open_meteo_pv_data_piece.py a test_data_preprocessing_piece.py. "
        "Žiadny test na LP, clipping, EMS ACK, OKTE parser, flexible loads, forecast-vs-actual. "
        "UC3.2: GitHub tests-dev.yml spúšťa pytest, ale v repositári nie sú test_*.py ani requirements-tests.txt. "
        "Bez testov každý zásah dvoch nových ľudí do dispatchu je riziko tichej regresie.",
    )

    add_para(doc, "PV „forecast“ sa učí zo syntetického labelu.", bold=True, space_after=4)
    add_para(
        doc,
        "OpenMeteoPVDataPiece počíta PVOUT z irradiance a performance ratio. "
        "To je slušný fallback, kým nie sú merania. Nie je to forecast z historickej výroby, "
        "ako to chce modul 5 referenčnej architektúry (historical RES production + weather + irradiance + asset). "
        "Kým nepríde meraný PV rad, ostáva to dočasný fallback — presne to, čo migrácia predpísala pre SolarSim.",
    )

    add_para(doc, "EMS Modbus je mapa, nie driver.", bold=True, space_after=4)
    add_para(
        doc,
        "build_register_map() vypíše CSV holding registers + coil. "
        "Do zariadenia sa netlačí. To je v poriadku, kým je cieľ REST EMS. "
        "Ak partner bude chcieť Modbus TCP, je to samostatná úloha, nie doriešenie existujúceho piece.",
    )

    add_heading(doc, "5.4 SEED / UC3.2 — čo ostalo po oddelení", 2)
    add_bullet(doc, "FeasibilityReportPiece v UC3.2 nie je. Dashboard a mrk_savings_report.json ho čiastočne zastupujú.")
    add_bullet(doc, "seed_handoff.json, ktorý spomína vyhodnotenie migrácie, v UC3.2 kóde nie je. SoMES dashboard má len seed_appendix poznámku („delegate_to: UC3.2_SEED“). Kontrakt medzi workflowmi nie je formalizovaný.")
    add_bullet(doc, "CatalogSync odkazuje na lokálny fallback katalog/, ktorý v repositári nie je.")
    add_bullet(doc, "OneData default token je v source (onedata_defaults.py) s komentárom odstrániť pred public release.")
    add_bullet(doc, "README odkazuje na PowerShell skripty (setup CI, delete pipelines), ktoré v strome nie sú.")
    add_bullet(doc, "SimulatePiece ostáva monolit: sizing, PV synth, battery, finance. SEED aj SoMES z neho historicky pijú. Refaktor je potrebný, ale je to riziková úloha — nie na prvý týždeň.")

    # ---------------------------------------------------------------------
    add_heading(doc, "6. Návrh práce pre dvoch programátorov", 1)
    add_para(
        doc,
        "Predpoklad: dvaja mid-level programátori (Python, pandas, REST, základy energetiky výhodou). "
        "Jeden architect / Filip ostáva owner architektúry a Domino importu. "
        "Noví ľudia nedotýkajú referenčných 18 modulov ako „návrhu“; dotýkajú sa implementácie a produkcie.",
    )

    add_heading(doc, "6.1 Programátor A — SoMES production a integrácie", 2)
    add_para(doc, "Profil: backend, integrácie, data contracts, observabilita. Denný repositár: scditech/uc3_3_somes.", italic=True)
    add_para(doc, "Balík A1 — Konektory na reálne zdroje", bold=True, space_after=4)
    add_bullet(doc, "Dohodnúť a zdokumentovať schému PV merania, BESS telemetrie, grid constraints, flexible loads (JSON/CSV + príklady).")
    add_bullet(doc, "Doplniť SomesConnectorsPiece: validácia schémy, jasný fail keď allow_demo_fallback=false, metriky provenance v dashboarde.")
    add_bullet(doc, "Odstrániť tiché demo fallback v produkčnom profile (demo ostane len pre CI a sandbox).")
    add_bullet(doc, "Acceptance: beh s allow_demo_fallback=false prejde, keď sú file/URL k dispozícii; bez nich spadne s čitateľnou chybou, nie so syntetikou.")

    add_para(doc, "Balík A2 — EMS/BEMS staging", bold=True, space_after=4)
    add_bullet(doc, "Mock EMS (malý FastAPI/flask endpoint) s ACK/reject pre CI.")
    add_bullet(doc, "End-to-end: approved plán → POST → ACK → záznam v dashboarde; reject plán sa neodošle (toto kód už robí, treba test a runbook).")
    add_bullet(doc, "Secrets mimo kódu (Domino secrets / env), žiadny token v defaults.")
    add_bullet(doc, "Acceptance: test s mock EMS; dokument mapovania polí pre partnera; require_delivery=true správa sa predvídateľne.")

    add_para(doc, "Balík A3 — Zatvorenie feedback slučky", bold=True, space_after=4)
    add_bullet(doc, "ModelMonitoring musí porovnávať ten istý plán, ktorý odišiel do EMS (node 114), nie preview na historickom loade.")
    add_bullet(doc, "Keď MAPE > prah, vygenerovať retraining dataset a buď spustiť IncrementalTrain, alebo aspoň one-click/artefakt, ktorý to spustí v ďalšom rune.")
    add_bullet(doc, "Dashboard: forecast vs actual, EMS delivery status, demo_generated flagy na dátových zdrojoch.")
    add_bullet(doc, "Acceptance: po nameranom dni je vidieť error metriky a jednoznačné next action; nie len JSON na disku.")

    add_para(doc, "Balík A4 — Prevádzkové testy konektorov a EMS", bold=True, space_after=4)
    add_bullet(doc, "Pytest: OKTE parser, Open-Meteo mapper, demo marker, EMS retry/ACK, schema connectors.")
    add_bullet(doc, "Zaradiť do CI tak, aby nepotrebovali live sieť (VCR/fixtures).")

    add_heading(doc, "6.2 Programátor B — optimalizácia, validácia, SEED kontrakt", 2)
    add_para(doc, "Profil: algoritmy, numerika, testy, čistý Python. Denné repositáre: uc3_3_somes + industry_sg_vre_cost_optimizer.", italic=True)
    add_para(doc, "Balík B1 — Test harness pre dispatch a validáciu (prvá úloha, povinná)", bold=True, space_after=4)
    add_bullet(doc, "Unit testy lp_battery_dispatch: energy balance, SOC band, complementarity charge/discharge, fallback greedy.")
    add_bullet(doc, "Unit testy GridFeasibility: C-rate, inverter, import/export, clipping korekčného plánu, approved flag.")
    add_bullet(doc, "Golden CSV na 96 krokov (jeden deň) — regresia objektívnej hodnoty a SOC trajektórie.")
    add_bullet(doc, "Acceptance: zmena v dispatch.py bez zmeny golden súboru musí spadnúť v CI. Toto je poistka pred príchodom ďalších úprav.")

    add_para(doc, "Balík B2 — Flexible loads do spoločnej optimalizácie", bold=True, space_after=4)
    add_bullet(doc, "Dnes: greedy window, potom pripočítať k loadu.")
    add_bullet(doc, "Cieľ: flexible load ako rozhodovacia premenná v LP, alebo aspoň iterácia „shift → dispatch → re-score“ s constraint max demand.")
    add_bullet(doc, "Nesmie rozbiť existujúci HiGHS model. Najprv testy (B1), potom rozšírenie.")
    add_bullet(doc, "Acceptance: dva flex loady s prekryvom okien nedostanú oba ten istý peak; connection headroom sa dodrží.")

    add_para(doc, "Balík B3 — Grid feasibility enhancement (až keď je site model)", bold=True, space_after=4)
    add_bullet(doc, "Nechať linearized PCC ako default (musí ostať, lebo site často nebude mať sieťový model).")
    add_bullet(doc, "Voliteľný mód: pandapower/simple radial feeder, ak partner dodá R/X, transformátor, limity napätia.")
    add_bullet(doc, "Kým nie sú dáta siete, toto je spike + interface, nie sľub full DMS.")
    add_bullet(doc, "Acceptance: pri absencii network modelu sa správa ako dnes; pri jeho prítomnosti vie povedať voltage violation na PCC aj na jednom internom node.")

    add_para(doc, "Balík B4 — SEED kontrakt a upratanie UC3.2", bold=True, space_after=4)
    add_bullet(doc, "Definovať seed_handoff.json (sizing kWp/kWh, limity batérie, inverter, connection) ako vstup SoMES scenario — nie CAPEX/NPV.")
    add_bullet(doc, "FeasibilityReportPiece v UC3.2, ak ho SEED naozaj potrebuje; inak explicitne povedať, že Dashboard + investment_evaluation.json stačia, a z dokumentácie FeasibilityReport vypustiť.")
    add_bullet(doc, "Odstrániť hardcoded OneData token; doplniť catalog/ fallback; opraviť tests-dev.yml (requirements-tests.txt + aspoň smoke testy pieces).")
    add_bullet(doc, "Nesiahnuť na SimulatePiece monolit, kým nie sú testy. Refaktor monolitu je follow-up po B1 na UC3.2 strane.")
    add_bullet(doc, "Acceptance: SoMES vie naštartovať D+1 plán zo SEED handoff bez toho, aby čítal payback. UC3.2 CI spustí pytest.")

    add_para(doc, "Balík B5 — Upratanie leftover pieces v SoMES", bold=True, space_after=4)
    add_bullet(doc, "S architektonickým ownerom rozhodnúť: archive / delete / optional DAG pre Evaluate a Explainable.")
    add_bullet(doc, "Zosúladiť README (stále spomína image 0.2.2) s config.toml 0.2.7.")
    add_bullet(doc, "Premenovať BatterySimPiece v UI labele na Dispatch / Battery Plan, aby Domino graf hovoril jazykom SoMES, nie UC3.2 simulácie.")

    add_heading(doc, "6.3 Čo im nedelegovať", 2)
    add_bullet(doc, "Prepis 18-modulovej architektúry alebo pridávanie investičných modulov do SoMES DAG.")
    add_bullet(doc, "Výmena HiGHS za iný solver bez merania na golden testoch.")
    add_bullet(doc, "Veľký refaktor SimulatePiece v UC3.2 pred existenciou testov.")
    add_bullet(doc, "Live Modbus TCP do zariadenia, kým partner nepovie, že REST nestačí.")
    add_bullet(doc, "Full distribution power-flow bez site modelu.")
    add_bullet(doc, "Domino CI/Harbor/GHCR plumbing — to je existujúca platforma, nie náborová úloha.")

    add_heading(doc, "6.4 Odporúčané poradie (bez kalendára)", 2)
    add_para(doc, "Spoločný onboarding oboch (kód, DAG, OneData, demo run). Potom paralelne:")
    add_bullet(doc, "A1 + B1 naraz. Integrácie a test harness sa nebijú a hneď znižujú riziko.")
    add_bullet(doc, "A2 + B2 potom. EMS staging a flex-load LP sú nezávislé.")
    add_bullet(doc, "A3 + B4. Feedback loop a SEED kontrakt, keď sú dáta a testy.")
    add_bullet(doc, "B3 a B5 na koniec, alebo keď príde site model / keď DAG treba zoštíhliť.")

    add_heading(doc, "6.5 Definition of Done pre náborový šprint", 2)
    add_bullet(doc, "SoMES demo profil ostáva, ale produkčný profil bez demo fallbacku je zdokumentovaný a testovaný.")
    add_bullet(doc, "Dispatch a GridFeasibility majú CI testy.")
    add_bullet(doc, "EMS má mock + runbook; reálny endpoint sa zapojí jedným secretom, nie zmenou kódu.")
    add_bullet(doc, "Je jasné, ktorý plán sa meria vo forecast-vs-actual (ten, ktorý odišiel do EMS).")
    add_bullet(doc, "SEED handoff je jeden JSON kontrakt, nie kopírovanie SimulatePiece do SoMES.")
    add_bullet(doc, "Leftover pieces sú buď v DAG so zmyslom, alebo označené ako non-production.")

    # ---------------------------------------------------------------------
    add_heading(doc, "7. Riziká, ak to nedelegujeme čisto", 1)
    add_bullet(doc, "Dvaja ľudia začnú „vylepšovať“ BatterySim aj Simulate naraz v oboch repositároch → divergencia dispatch logiky.")
    add_bullet(doc, "Niekto vráti InvestmentEval do SoMES DAG, lebo dashboard ešte vie zobraziť payback appendix.")
    add_bullet(doc, "Zapne sa allow_demo_fallback v produkcii a reporty pôjdu von so syntetickým BESS SOC.")
    add_bullet(doc, "Grid power-flow sa sľúbi zákazníkovi skôr, než existuje sieťový model.")
    add_bullet(doc, "Feedback loop sa označí za closed, hoci IncrementalTrain treba spúšťať ručne.")

    # ---------------------------------------------------------------------
    add_heading(doc, "8. Wording do M24 / D3.3 (odporúčanie)", 1)
    add_para(
        doc,
        "Nasledujúci odsek je v súlade so zadaním aj s kódom. "
        "Jediná vec, ktorú kód núti spresniť: remaining items nie sú len externé závislosti, "
        "ale aj produkčné doriešenie testov, closed-loop retraining a voliteľný power-flow.",
        italic=True,
    )
    add_para(
        doc,
        "The migration of the original UC3.3 Sustainable Model into the SoMES workflow family "
        "can be considered successfully completed from an architectural perspective. The resulting "
        "SoMES implementation fully adopts the operational focus defined by the SoMES Workflow Type "
        "Reference Architecture, including forecasting, dispatch optimisation, battery scheduling, "
        "grid import/export planning, technical validation, EMS/BEMS integration, KPI monitoring, "
        "dashboarding and continuous feedback-driven learning. Investment-oriented functionality has "
        "been successfully separated from the core workflow and transferred to the SEED domain. "
        "Remaining open items are limited to external integration dependencies (real PV measurements, "
        "BESS telemetry, EMS endpoint connectivity), production hardening of the feedback loop and "
        "test coverage, and optional future enhancement of grid validation through full network "
        "power-flow modelling.",
    )

    # ---------------------------------------------------------------------
    add_heading(doc, "9. Príloha — inventár pieces", 1)

    add_heading(doc, "9.1 SoMES (scditech/uc3_3_somes) — v Domino importe", 2)
    add_table(
        doc,
        ["Piece", "Logický SoMES modul"],
        [
            ["FetchEnergyDataPiece", "1 Ingestion"],
            ["PreprocessEnergyDataPiece", "2–3 Validation / Preprocessing"],
            ["TrainModelPiece / PredictPiece / ForecastHorizonPiece / IncrementalTrainPiece", "4 Load Forecast"],
            ["OpenMeteoPVDataPiece + PVOUT chain + PvoutToVirtualSolarPiece", "5 PV / RES Forecast"],
            ["SomesConnectorsPiece", "1 Ingestion (enabling)"],
            ["PriceForecastPiece", "6 Price Forecast (enabling)"],
            ["BatteryStrategyOptimizerPiece", "7–8 prahy pre greedy fallback"],
            ["BatterySimPiece", "7 Dispatch + 8 Battery plan + 9 Import/export"],
            ["FlexibleLoadSchedulePiece", "10 Flexible loads"],
            ["GridFeasibilityPiece", "11–13 Technical validation"],
            ["EmsBemsOutputPiece", "14 Next-day plan doručenie + 17 EMS interface"],
            ["DashboardPiece", "15 KPI vizualizácia + 16 Dashboard"],
            ["ModelMonitoringPiece", "18 Forecast vs actual"],
            ["DataPreprocessing / DataNormalization / ModelDecider / Pvout* / Inference", "PV vetva (UC3.4 enabling)"],
        ],
    )

    add_heading(doc, "9.2 UC3.2 (industry_sg_vre_cost_optimizer)", 2)
    add_table(
        doc,
        ["Piece", "SEED / overlap"],
        [
            ["UserInputPiece", "SEED vstup"],
            ["CatalogSyncPiece / CatalogRankerPiece", "SEED — hardvér"],
            ["TechnicalLimitsPiece / SizingOptimizationPiece", "SEED — sizing (z SoMES odstránené)"],
            ["SolarSimPiece / BatteryStrategyOptimizerPiece / BatterySimPiece / SimulatePiece", "Overlap s SoMES — treba kontrakt, nie duplicitný rozvoj"],
            ["KPIPiece / InvestmentEvalPiece / DashboardPiece", "SEED — investičné KPI / CFO"],
            ["FeasibilityReportPiece", "Chýba"],
        ],
    )

    add_heading(doc, "10. Záver pre rozhodnutie o nábore", 1)
    add_para(
        doc,
        "SoMES netreba stavať znova. Treba ho dostať z architektonicky hotového stavu do prevádzkovo tvrdého. "
        "Dvaja programátori to zvládnu, ak jeden drží integrácie a druhý testy/algoritmy/SEED kontrakt, "
        "a ak im architect zakáže vracať investičnú logiku do operačného DAG. "
        "Najväčšia hodnota náboru nie je „ďalší forecast model“, ale: reálne dáta, EMS na ostro, "
        "testy dispatchu a jeden čistý handoff medzi SEED a SoMES.",
    )
    add_para(
        doc,
        "Dokument overený voči verejnému kódu scditech/uc3_3_somes @ 0.2.7 (posledný push 15. 8. 2026) "
        "a filipchrvala/industry_sg_vre_cost_optimizer @ 0.1.42 (posledný obsahový commit 28. 6. 2026).",
        italic=True,
        size=10,
    )

    out = Path("/workspace/docs/UC3.3_SoMES_gap_analyza_delegovanie_prace.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


if __name__ == "__main__":
    path = build()
    print(path)
    print("bytes", path.stat().st_size)
